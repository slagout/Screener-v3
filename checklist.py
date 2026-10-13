import io
import os
from datetime import datetime

import openpyxl
from openpyxl.cell.cell import MergedCell

TEMPLATE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "CAGR_Research_Checklist_v4.1.xlsx")
MASTER_SHEET = "Cagr Trade Checklist Master"

# Judgement cells that ship with sample answers in the template; blank them so stale picks don't leak in.
MANUAL_CELLS = [
    "C31", "C46", "C61", "C75", "C79", "C124", "C126", "C128", "C130", "C146", "C173",
    "C235", "C247", "C251", "C267", "C271", "C290", "C317", "C322", "C324", "C329",
    "C343", "C358", "C159", "C42", "C5",
]


def _row(df, names):
    if df is None or getattr(df, "empty", True):
        return None
    for n in names:
        if n in df.index:
            return df.loc[n]
    return None


def _num(x):
    try:
        v = float(x)
        return None if v != v else v
    except (TypeError, ValueError):
        return None


def _pct(x):
    v = _num(x)
    return None if v is None else round(v * 100, 2)


def _safe(fn, default=None):
    try:
        return fn()
    except Exception:
        return default


def _altman_z(stock, info, fin, bal):
    ta = _num(_row(bal, ["Total Assets"]).iloc[0]) if _row(bal, ["Total Assets"]) is not None else None
    tl_row = _row(bal, ["Total Liabilities Net Minority Interest", "Total Liabilities"])
    tl = _num(tl_row.iloc[0]) if tl_row is not None else None
    wc_row = _row(bal, ["Working Capital"])
    re_row = _row(bal, ["Retained Earnings"])
    ebit_row = _row(fin, ["EBIT", "Operating Income"])
    rev_row = _row(fin, ["Total Revenue", "Revenue"])
    mcap = _num(info.get("marketCap"))
    if None in (ta, tl) or not ta or not tl or mcap is None or wc_row is None or re_row is None \
            or ebit_row is None or rev_row is None:
        return None
    wc, re_, ebit, rev = (_num(wc_row.iloc[0]), _num(re_row.iloc[0]),
                          _num(ebit_row.iloc[0]), _num(rev_row.iloc[0]))
    if None in (wc, re_, ebit, rev):
        return None
    return round(1.2 * wc / ta + 1.4 * re_ / ta + 3.3 * ebit / ta + 0.6 * mcap / tl + rev / ta, 2)


def build_fields(stock, info, q):
    """Return {cell: (value, source, label, section)} for the master sheet. Missing data is skipped."""
    out = {}

    def put(cell, value, source, label, section):
        if value is not None and value != "":
            out[cell] = (value, source, label, section)

    fin = _safe(lambda: stock.financials)
    bal = _safe(lambda: stock.balance_sheet)
    cf = _safe(lambda: stock.cashflow)

    put("B2", q.get("company_name"), "Screener", "Company", "Header")
    put("D2", q.get("ticker"), "Screener", "Ticker", "Header")
    put("B3", datetime.now().strftime("%Y-%m-%d"), "Today", "Date", "Header")
    put("C5", "Yes", "Screener price range", "Price in range with portfolio allocation?", "Header")
    put("A10", "Found through the Francine Stock Screener (options-liquid, positive margin, IV in range).",
        "Screener", "How did you become aware of this company?", "Section 1")
    put("A18", info.get("longBusinessSummary"), "Yahoo Finance", "Short description", "Section 2")

    ipo_ms = _num(info.get("firstTradeDateMilliseconds") or info.get("firstTradeDateEpochUtc"))
    if ipo_ms:
        ts = ipo_ms / 1000 if ipo_ms > 1e11 else ipo_ms
        put("C40", datetime.utcfromtimestamp(ts).year, "Yahoo Finance", "Year it went public (first trade)", "Section 3")
    put("C42", "Yes" if q.get("options_liquid") else None, "Screener", "Weekly options available?", "Section 3")
    put("C44", q.get("iv"), "Screener (IV %)", "Average volatility (IV %)", "Section 3")

    revs = []
    rev_row = _row(fin, ["Total Revenue", "Revenue"])
    if rev_row is not None:
        revs = [_num(v) for v in rev_row.tolist()]
        revs = [v for v in revs if v]
    yoy = None
    if len(revs) >= 2 and revs[1] > 0:
        yoy = round((revs[0] / revs[1] - 1) * 100, 2)
    if len(revs) >= 3:
        growing = all(revs[i] > revs[i + 1] for i in range(len(revs) - 1))
        put("C151", "Yes" if growing else "No", "Yahoo Finance financials", "Is revenue growing steadily?", "Section 7")
    put("D151", f"{yoy} % year over year" if yoy is not None else None, "Yahoo Finance financials",
        "Latest revenue growth (YoY)", "Section 7")

    ni_row = _row(fin, ["Net Income"])
    ni = [_num(v) for v in ni_row.tolist()] if ni_row is not None else []
    ni = [v for v in ni if v is not None]
    if len(ni) >= 3:
        put("C153", "Yes" if all(ni[i] > ni[i + 1] for i in range(len(ni) - 1)) else "No",
            "Yahoo Finance financials", "Are earnings growing steadily?", "Section 7")
    if len(ni) >= 3 and ni[-1] > 0 and ni[0] > 0:
        egr = round(((ni[0] / ni[-1]) ** (1 / (len(ni) - 1)) - 1) * 100, 2)
        put("C211", egr, "Yahoo Finance financials", "Earnings growth (CAGR %)", "Section 9")

    put("C59", q.get("cagr"), "Screener revenue CAGR", "Revenue growth (3Yr CAGR %)", "Section 4")
    put("C209", q.get("cagr"), "Screener revenue CAGR", "Revenue growth (%)", "Section 9")

    fwd = _safe(lambda: stock.growth_estimates)
    if fwd is not None and not fwd.empty:
        col = "stockTrend" if "stockTrend" in fwd.columns else fwd.columns[0]
        for k in ("+1y", "+5y", "0y"):
            if k in fwd.index and _num(fwd.loc[k, col]) is not None:
                put("C63", round(_num(fwd.loc[k, col]) * 100, 2), "Yahoo Finance estimates",
                    "Forward growth (%)", "Section 4")
                break

    debt, cash, fcf = _num(info.get("totalDebt")), _num(info.get("totalCash")), _num(info.get("freeCashflow"))
    if debt is not None and cash is not None and fcf:
        put("C65", round((debt - cash) / fcf, 2), "Yahoo Finance", "Net debt to FCF ratio", "Section 4")
    if debt is not None and fcf and fcf > 0:
        put("C67", round(debt / fcf, 1), "Yahoo Finance", "Years to repay debt (debt / FCF)", "Section 4")
    put("C69", _safe(lambda: _altman_z(stock, info, fin, bal)), "Computed from statements",
        "Altman-Z score (3+ required)", "Section 4")

    # Buybacks: shares outstanding falling across the balance sheet history.
    sh_row = _row(bal, ["Ordinary Shares Number", "Share Issued"])
    if sh_row is not None:
        sh = [_num(v) for v in sh_row.tolist() if _num(v)]
        if len(sh) >= 2:
            put("C155", "Yes" if sh[0] < sh[-1] else "No", "Yahoo Finance balance sheet",
                "Stock buybacks (shares reducing)?", "Section 7")
            put("C261", "Yes" if sh[0] < sh[-1] else "No", "Yahoo Finance balance sheet",
                "Do they repurchase shares?", "Section 10")

    div_rate = _num(info.get("dividendRate")) or _num(info.get("trailingAnnualDividendRate"))
    pays = bool(div_rate and div_rate > 0)
    put("C157", "Yes" if pays else "No", "Yahoo Finance", "Dividends?", "Section 7")
    put("C263", "Yes" if pays else "No", "Yahoo Finance", "Do they pay dividends?", "Section 10")
    if pays:
        dy = _num(info.get("dividendYield"))
        if dy is not None:
            put("D157", f"{dy if dy > 1 else round(dy * 100, 2)} %", "Yahoo Finance", "Dividend yield", "Section 7")

    mcap = _num(info.get("marketCap"))
    if cf is not None and mcap:
        parts = []
        for names in (["Cash Dividends Paid"], ["Repurchase Of Capital Stock"], ["Repayment Of Debt"]):
            r = _row(cf, names)
            v = _num(r.iloc[0]) if r is not None else None
            parts.append(-v if v and v < 0 else 0)
        put("C161", round(sum(parts) / mcap * 100, 2), "Computed from cash flow",
            "Shareholder yield % (div + buyback + debt paydown)", "Section 7")

    officers = info.get("companyOfficers") or []
    ceo = next((o for o in officers if "ceo" in (o.get("title") or "").lower()
                or "chief executive" in (o.get("title") or "").lower()), None)
    if ceo:
        put("C239", ceo.get("name"), "Yahoo Finance", "Who is the CEO?", "Section 10")
    insiders = _pct(info.get("heldPercentInsiders"))
    put("C243", f"{insiders}% (all insiders)" if insiders is not None else None, "Yahoo Finance",
        "How much stock do they own?", "Section 10")

    ed = _safe(lambda: stock.get_earnings_dates(limit=12))
    if ed is not None and not ed.empty and "Surprise(%)" in ed.columns:
        past = ed.dropna(subset=["Surprise(%)"]).head(8)
        if len(past):
            beats = int((past["Surprise(%)"] > 0).sum())
            put("C255", beats, f"Yahoo Finance (last {len(past)} quarters)", "Earnings beats out of 8", "Section 10")
            put("C253", "Yes" if beats >= 6 else ("Mixed" if beats >= 4 else "No"),
                "Yahoo Finance", "Beat analyst earnings expectations?", "Section 10")

    roe = _pct(info.get("returnOnEquity"))
    put("D259", f"ROE {roe}%" if roe is not None else None, "Yahoo Finance", "ROIC / FCF margin note", "Section 10")

    key = (info.get("recommendationKey") or "").lower()
    rating = {"strong_buy": "Strong Buy", "buy": "Buy", "hold": "Hold",
              "underperform": "Sell", "sell": "Sell", "strong_sell": "Sell"}.get(key)
    put("C333", rating, "Yahoo Finance analysts", "Analyst rating", "Section 12")
    put("B336", _num(info.get("targetLowPrice")), "Yahoo Finance analysts", "Price target low ($)", "Section 12")
    put("C336", _num(info.get("targetMeanPrice")), "Yahoo Finance analysts", "Price target avg ($)", "Section 12")
    put("D336", _num(info.get("targetHighPrice")), "Yahoo Finance analysts", "Price target high ($)", "Section 12")
    put("C339", q.get("price"), "Screener", "Current price ($)", "Section 12")
    return out


def build_xlsx(values):
    """values: {cell: value}. Fills the master sheet of the template and returns xlsx bytes."""
    wb = openpyxl.load_workbook(TEMPLATE)
    ws = wb[MASTER_SHEET]
    for sheet in list(wb.sheetnames):
        if sheet != MASTER_SHEET:
            del wb[sheet]

    def target(coord):
        c = ws[coord]
        if isinstance(c, MergedCell):
            for rng in ws.merged_cells.ranges:
                if coord in rng:
                    return ws.cell(rng.min_row, rng.min_col)
        return c

    for coord in MANUAL_CELLS:
        target(coord).value = None
    for coord, v in values.items():
        c = target(coord)
        if isinstance(v, (int, float)) and "%" in (c.number_format or ""):
            v = v / 100
        c.value = v
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
